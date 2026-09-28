"""Office content that lives outside the plain cell or paragraph text must not be dropped.

Every document here is generated at test time, with ``python-docx`` or ``openpyxl``.
Parts those libraries cannot author (footnotes, endnotes, a hyperlink with a missing
relationship, a formula's stored result) are added by rewriting the generated
package's XML.
"""

from __future__ import annotations

import json
import re
from collections.abc import Callable
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from typer.testing import CliRunner

import archiv.ingestion.normalize_docx as normalize_docx_module
import archiv.ingestion.normalize_xlsx as normalize_xlsx_module
from archiv.cli import app
from archiv.contracts import NormalizedDocument
from archiv.ingestion.normalizers import MalformedInputError, normalize

DIGEST = "0" * 64
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _save(document: DocxDocument, path: Path) -> Path:
    document.save(str(path))
    return path


def _by_locator(result: NormalizedDocument) -> dict[str, str]:
    return {
        json.dumps(segment.locator, sort_keys=True): segment.text for segment in result.segments
    }


def _key(**locator: object) -> str:
    return json.dumps(locator, sort_keys=True)


def _add_hyperlink(paragraph: Paragraph, url: str, text: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = parse_xml(
        f'<w:hyperlink xmlns:w="{W}" xmlns:r="{R}" r:id="{relationship_id}">'
        f"<w:r><w:t>{text}</w:t></w:r></w:hyperlink>"
    )
    paragraph._p.append(link)  # pyright: ignore[reportPrivateUsage, reportUnknownMemberType]


def _add_anchor_link(paragraph: Paragraph, anchor: str, text: str) -> None:
    link = parse_xml(
        f'<w:hyperlink xmlns:w="{W}" w:anchor="{anchor}"><w:r><w:t>{text}</w:t></w:r></w:hyperlink>'
    )
    paragraph._p.append(link)  # pyright: ignore[reportPrivateUsage, reportUnknownMemberType]


def _rewrite(
    path: Path, edits: dict[str, bytes], transform: dict[str, Callable[[str], str]]
) -> None:
    """Rewrite members of a generated package: replace ``edits``, apply ``transform``."""

    with ZipFile(path) as source:
        members = {info.filename: source.read(info.filename) for info in source.infolist()}
    for name, function in transform.items():
        members[name] = function(members[name].decode("utf-8")).encode("utf-8")
    members.update(edits)
    raw = BytesIO()
    with ZipFile(raw, "w", ZIP_DEFLATED) as target:
        for name, data in members.items():
            target.writestr(name, data)
    path.write_bytes(raw.getvalue())


def _note_part(tag: str, notes: dict[int, str]) -> bytes:
    """A footnotes or endnotes part with Word's two separator entries plus ``notes``."""

    entries = [
        f'<w:{tag} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{tag}>',
        f'<w:{tag} w:type="continuationSeparator" w:id="0"><w:p><w:r>'
        f"<w:continuationSeparator/></w:r></w:p></w:{tag}>",
    ]
    for note_id, text in notes.items():
        entries.append(
            f'<w:{tag} w:id="{note_id}"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:{tag}>'
        )
    return (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:{tag}s xmlns:w="{W}" xmlns:r="{R}">{"".join(entries)}</w:{tag}s>'
    ).encode()


def _add_notes(path: Path, *, footnotes: dict[int, str], endnotes: dict[int, str]) -> None:
    def relationships(xml: str) -> str:
        extra = (
            f'<Relationship Id="rIdFoot" Type="{RT.FOOTNOTES}" Target="footnotes.xml"/>'
            f'<Relationship Id="rIdEnd" Type="{RT.ENDNOTES}" Target="endnotes.xml"/>'
        )
        return xml.replace("</Relationships>", extra + "</Relationships>")

    def content_types(xml: str) -> str:
        base = "application/vnd.openxmlformats-officedocument.wordprocessingml"
        extra = (
            '<Override PartName="/word/footnotes.xml" '
            f'ContentType="{base}.footnotes+xml"/>'
            '<Override PartName="/word/endnotes.xml" '
            f'ContentType="{base}.endnotes+xml"/>'
        )
        return xml.replace("</Types>", extra + "</Types>")

    _rewrite(
        path,
        {
            "word/footnotes.xml": _note_part("footnote", footnotes),
            "word/endnotes.xml": _note_part("endnote", endnotes),
        },
        {
            "word/_rels/document.xml.rels": relationships,
            "[Content_Types].xml": content_types,
        },
    )


def test_docx_table_cells_are_extracted(tmp_path: Path) -> None:
    """The reproduction from the step: a paragraph, a table of names and dates."""

    document = Document()
    document.add_paragraph("Staff register")
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Joined"
    table.cell(0, 2).text = "Office"
    table.cell(1, 0).text = "Amina Khan"
    table.cell(1, 1).text = "2021-03-04"
    table.cell(1, 2).text = "Lahore"
    # A horizontal merge across two columns keeps the first column's locator.
    merged = table.cell(2, 0).merge(table.cell(2, 1))
    merged.text = "Pending review"
    table.cell(2, 2).text = ""
    document.add_paragraph("After the table")
    path = _save(document, tmp_path / "register.docx")

    result = normalize(path, DIGEST)
    located = _by_locator(result)

    assert located[_key(paragraph=1)] == "Staff register"
    # The table sits between two body paragraphs but does not shift their numbers.
    assert located[_key(paragraph=2)] == "After the table"
    assert located[_key(table=1, row=1, column=1)] == "Name"
    assert located[_key(table=1, row=2, column=1)] == "Amina Khan"
    assert located[_key(table=1, row=2, column=2)] == "2021-03-04"
    assert located[_key(table=1, row=2, column=3)] == "Lahore"
    assert located[_key(table=1, row=3, column=1)] == "Pending review"
    assert _key(table=1, row=3, column=2) not in located, "merged cell reported twice"
    assert _key(table=1, row=3, column=3) not in located, "empty cell produced a segment"

    assert len(result.tables) == 1
    assert result.tables[0].locator == {"table": 1}
    assert result.tables[0].rows == [
        ["Name", "Joined", "Office"],
        ["Amina Khan", "2021-03-04", "Lahore"],
        ["Pending review", None, None],
    ]
    assert result.metadata["paragraphs"] == 2
    assert result.metadata["tables"] == 1


def test_docx_nested_table_and_vertical_merge_keep_every_word(tmp_path: Path) -> None:
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "Spans two rows"
    table.cell(0, 1).text = "Outer"
    inner = table.cell(1, 1).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "Inside a nested table"
    path = _save(document, tmp_path / "nested.docx")

    located = _by_locator(normalize(path, DIGEST))

    assert located[_key(table=1, row=1, column=1)] == "Spans two rows"
    assert _key(table=1, row=2, column=1) not in located, "vertical merge reported twice"
    assert "Inside a nested table" in located[_key(table=1, row=2, column=2)]


def test_docx_headers_footers_and_footnotes_are_extracted(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Body text")
    first = document.sections[0]
    first.header.paragraphs[0].text = "Confidential header"
    first.footer.paragraphs[0].text = "Page footer text"
    first.different_first_page_header_footer = True
    first.first_page_header.paragraphs[0].text = "Cover page header"
    footer_table = first.footer.add_table(rows=1, cols=2, width=Inches(6))
    footer_table.cell(0, 1).text = "Footer cell"
    # A second section that inherits the header must not repeat it.
    document.add_section()
    document.add_paragraph("Second section body")
    path = _save(document, tmp_path / "stories.docx")
    _add_notes(
        path,
        footnotes={1: "Source: archive ledger 1998"},
        endnotes={1: "Closing endnote"},
    )

    result = normalize(path, DIGEST)
    located = _by_locator(result)

    assert located[_key(section=1, header="default", paragraph=1)] == "Confidential header"
    assert located[_key(section=1, footer="default", paragraph=1)] == "Page footer text"
    assert located[_key(section=1, header="first", paragraph=1)] == "Cover page header"
    assert located[_key(section=1, footer="default", table=1, row=1, column=2)] == "Footer cell"
    assert not any('"section": 2' in key for key in located), "linked header was repeated"
    assert located[_key(footnote=1, paragraph=1)] == "Source: archive ledger 1998"
    assert located[_key(endnote=1, paragraph=1)] == "Closing endnote"
    # Word's separator entries are layout, not content.
    assert not any('"footnote": -1' in key or '"footnote": 0' in key for key in located)
    assert [segment.text for segment in result.segments].count("Confidential header") == 1


def test_docx_hyperlink_targets_are_preserved(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("Read the ")
    _add_hyperlink(paragraph, "https://example.org/policy", "policy")
    paragraph.add_run(" and the ")
    _add_anchor_link(paragraph, "appendix_b", "appendix")
    cell_paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    _add_hyperlink(cell_paragraph, "https://example.org/cell", "cell link")
    path = _save(document, tmp_path / "links.docx")

    result = normalize(path, DIGEST)
    located = _by_locator(result)

    # The visible link text stays in the paragraph, as before.
    assert located[_key(paragraph=1)] == "Read the policy and the appendix"
    # The targets, which were silently lost, now each have a segment.
    assert located[_key(paragraph=1, hyperlink=1)] == "https://example.org/policy"
    assert located[_key(paragraph=1, hyperlink=2)] == "#appendix_b"
    assert located[_key(table=1, row=1, column=1)] == "cell link"
    assert located[_key(table=1, row=1, column=1, hyperlink=1)] == "https://example.org/cell"


def test_docx_hyperlink_with_missing_relationship_is_skipped_not_guessed(
    tmp_path: Path,
) -> None:
    document = Document()
    paragraph = document.add_paragraph("See ")
    _add_hyperlink(paragraph, "https://example.org/gone", "here")
    path = _save(document, tmp_path / "broken-link.docx")
    _rewrite(
        path,
        {},
        {
            "word/_rels/document.xml.rels": lambda xml: re.sub(
                r'<Relationship [^>]*Target="https://example.org/gone"[^>]*/>', "", xml
            )
        },
    )

    located = _by_locator(normalize(path, DIGEST))

    assert located[_key(paragraph=1)] == "See here"
    assert not any("hyperlink" in key for key in located)


def test_docx_comments_are_extracted(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("Disputed figure: 4,200")
    document.add_comment(paragraph.runs, text="Check against the 2019 audit", author="Reviewer")
    path = _save(document, tmp_path / "comments.docx")

    located = _by_locator(normalize(path, DIGEST))

    assert located[_key(paragraph=1)] == "Disputed figure: 4,200"
    comment_keys = [key for key in located if '"comment"' in key]
    assert len(comment_keys) == 1
    assert located[comment_keys[0]] == "Check against the 2019 audit"


def test_docx_heading_levels_are_recorded(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("Findings", level=1)
    document.add_paragraph("Plain body")
    document.add_heading("Method", level=2)
    path = _save(document, tmp_path / "headings.docx")

    located = _by_locator(normalize(path, DIGEST))

    assert located[_key(paragraph=1, heading_level=1)] == "Findings"
    assert located[_key(paragraph=2)] == "Plain body"
    assert located[_key(paragraph=3, heading_level=2)] == "Method"


def test_docx_body_text_and_numbering_are_unchanged_from_python_docx(tmp_path: Path) -> None:
    """Citations made by the paragraphs-only reader must still point at the same text."""

    document = Document()
    document.add_paragraph("First")
    document.add_paragraph("")
    run_paragraph = document.add_paragraph("Tab\tand")
    run_paragraph.add_run().add_break()
    run_paragraph.add_run("line break")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "cell"
    document.add_paragraph("Last")
    path = _save(document, tmp_path / "compat.docx")

    result = normalize(path, DIGEST)
    reopened = Document(str(path))
    expected = {
        _key(paragraph=index): paragraph.text
        for index, paragraph in enumerate(reopened.paragraphs, 1)
        if paragraph.text
    }

    body = {key: text for key, text in _by_locator(result).items() if '"table"' not in key}
    assert body == expected
    assert result.metadata["paragraphs"] == len(reopened.paragraphs)


def test_docx_table_cell_limit_refuses_the_document(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    document = Document()
    document.add_table(rows=2, cols=2)
    path = _save(document, tmp_path / "large.docx")
    monkeypatch.setattr(normalize_docx_module, "MAX_TABLE_CELLS", 3)

    with pytest.raises(MalformedInputError, match="table cell limit"):
        normalize(path, DIGEST)


def test_docx_wide_column_span_counts_against_the_table_budget(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """One cell claiming a huge span must not build a huge grid from a tiny file."""

    document = Document()
    table = document.add_table(rows=3, cols=1)
    for row in table.rows:
        row.cells[0].text = "x"
    path = _save(document, tmp_path / "wide.docx")
    _rewrite(
        path,
        {},
        {
            "word/document.xml": lambda xml: xml.replace(
                "<w:tcPr>", '<w:tcPr><w:gridSpan w:val="900"/>', 1
            )
        },
    )
    # Three cells, but 900 + 900 + 900 grid positions once the rows are padded.
    monkeypatch.setattr(normalize_docx_module, "MAX_TABLE_CELLS", 2_000)

    with pytest.raises(MalformedInputError, match="table cell limit"):
        normalize(path, DIGEST)


def test_docx_span_beyond_the_column_ceiling_refuses_the_document(tmp_path: Path) -> None:
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "x"
    path = _save(document, tmp_path / "too-wide.docx")
    _rewrite(
        path,
        {},
        {
            "word/document.xml": lambda xml: xml.replace(
                "<w:tcPr>", '<w:tcPr><w:gridSpan w:val="200000"/>', 1
            )
        },
    )

    with pytest.raises(MalformedInputError, match="gridSpan"):
        normalize(path, DIGEST)


def test_docx_entity_declaration_in_footnotes_refuses_the_document(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Body")
    path = _save(document, tmp_path / "entity.docx")
    _add_notes(path, footnotes={1: "&big;"}, endnotes={})
    with ZipFile(path) as archive:
        footnotes = archive.read("word/footnotes.xml").decode("utf-8")
    declaration = '<!DOCTYPE w:footnotes [<!ENTITY big "' + "A" * 64 + '">]>'
    footnotes = footnotes.replace("?>", "?>" + declaration, 1)
    _rewrite(path, {"word/footnotes.xml": footnotes.encode("utf-8")}, {})

    with pytest.raises(MalformedInputError, match="declarations and entities"):
        normalize(path, DIGEST)


def test_docx_entity_declaration_in_utf16_footnotes_refuses_the_document(
    tmp_path: Path,
) -> None:
    """A byte search misses UTF-16; the refusal must not depend on the encoding."""

    document = Document()
    document.add_paragraph("Body")
    path = _save(document, tmp_path / "entity-utf16.docx")
    _add_notes(path, footnotes={1: "&big;"}, endnotes={})
    with ZipFile(path) as archive:
        footnotes = archive.read("word/footnotes.xml").decode("utf-8")
    declaration = '<!DOCTYPE w:footnotes [<!ENTITY big "' + "A" * 64 + '">]>'
    footnotes = footnotes.replace('encoding="UTF-8"', 'encoding="UTF-16"', 1)
    footnotes = footnotes.replace("?>", "?>" + declaration, 1)
    encoded = footnotes.encode("utf-16")
    assert b"<!DOCTYPE" not in encoded.upper()
    _rewrite(path, {"word/footnotes.xml": encoded}, {})

    with pytest.raises(MalformedInputError, match="declarations and entities"):
        normalize(path, DIGEST)


def test_docx_malformed_note_identifier_is_kept_not_refused(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Body")
    path = _save(document, tmp_path / "note-id.docx")
    _add_notes(path, footnotes={1: "Odd note"}, endnotes={})
    with ZipFile(path) as archive:
        footnotes = archive.read("word/footnotes.xml").decode("utf-8")
    _rewrite(
        path,
        {"word/footnotes.xml": footnotes.replace('w:id="1"', 'w:id="--5"').encode("utf-8")},
        {},
    )

    located = _by_locator(normalize(path, DIGEST))

    assert located[_key(paragraph=1)] == "Body"
    assert located[_key(footnote="--5", paragraph=1)] == "Odd note"


def test_docx_outline_level_nine_marks_body_text_under_a_heading_style(tmp_path: Path) -> None:
    document = Document()
    heading = document.add_heading("Styled as a heading", level=1)
    properties = heading._p.get_or_add_pPr()  # pyright: ignore[reportPrivateUsage]
    properties.append(parse_xml(f'<w:outlineLvl xmlns:w="{W}" w:val="9"/>'))  # pyright: ignore[reportUnknownMemberType]
    path = _save(document, tmp_path / "outline.docx")

    located = _by_locator(normalize(path, DIGEST))

    assert located[_key(paragraph=1)] == "Styled as a heading"


def test_docx_invalid_grid_span_refuses_the_document(tmp_path: Path) -> None:
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "x"
    path = _save(document, tmp_path / "span.docx")
    _rewrite(
        path,
        {},
        {
            "word/document.xml": lambda xml: xml.replace(
                "<w:tcPr>", '<w:tcPr><w:gridSpan w:val="two"/>', 1
            )
        },
    )

    with pytest.raises(MalformedInputError, match="gridSpan"):
        normalize(path, DIGEST)


def test_docx_table_cell_is_found_by_search_with_its_table_locator(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Register")
    document.add_table(rows=1, cols=2).cell(0, 1).text = "ARCHIV-TABLE-MARKER-2026"
    path = _save(document, tmp_path / "cited.docx")
    home = tmp_path / "home"
    runner = CliRunner()

    ingested = runner.invoke(app, ["ingest", str(path), "--home", str(home)])
    assert ingested.exit_code == 0, ingested.output
    rebuilt = runner.invoke(app, ["rebuild-search-index", "--home", str(home)])
    assert rebuilt.exit_code == 0, rebuilt.output
    found = runner.invoke(app, ["search", "ARCHIV-TABLE-MARKER-2026", "--home", str(home)])

    assert found.exit_code == 0, found.output
    payload = json.loads(found.output)
    assert [hit["citation"]["locator"] for hit in payload] == [{"table": 1, "row": 1, "column": 2}]


# --- Excel workbooks -----------------------------------------------------------------
#
# ``openpyxl`` writes a formula with an empty result, as if never calculated. A file
# saved by a spreadsheet application stores the last calculated result beside the
# formula, so the tests write that result into the generated package's XML.

XLSX_SHEET = "xl/worksheets/sheet1.xml"


def _save_workbook(workbook: Workbook, path: Path) -> Path:
    workbook.save(str(path))
    return path


def _store_result(path: Path, formula: str, result: str, *, part: str = XLSX_SHEET) -> None:
    written = f"<f>{formula}</f><v></v>"

    def store(xml: str) -> str:
        assert written in xml, xml
        return xml.replace(written, f"<f>{formula}</f><v>{result}</v>")

    _rewrite(path, {}, {part: store})


def _budget_workbook() -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Budget"
    sheet["A2"] = "Salaries"
    sheet["B2"] = 750000
    sheet["A3"] = "Rent"
    sheet["B3"] = 500000
    sheet["A4"] = "Total"
    sheet["B4"] = "=SUM(B2:B3)"
    return workbook


def _search(path: Path, home: Path, query: str) -> list[dict[str, object]]:
    runner = CliRunner()
    ingested = runner.invoke(app, ["ingest", str(path), "--home", str(home)])
    assert ingested.exit_code == 0, ingested.output
    rebuilt = runner.invoke(app, ["rebuild-search-index", "--home", str(home)])
    assert rebuilt.exit_code == 0, rebuilt.output
    found = runner.invoke(app, ["search", query, "--home", str(home)])
    assert found.exit_code == 0, found.output
    return json.loads(found.output)


def test_formula_cell_indexes_the_computed_value(tmp_path: Path) -> None:
    path = _save_workbook(_budget_workbook(), tmp_path / "budget.xlsx")
    _store_result(path, "SUM(B2:B3)", "1250000")

    located = _by_locator(normalize(path, DIGEST))

    assert located[_key(sheet="Budget", cell="B4", formula="=SUM(B2:B3)")] == "1250000"
    assert "=SUM(B2:B3)" not in located.values()


def test_formula_text_is_preserved_in_the_locator(tmp_path: Path) -> None:
    path = _save_workbook(_budget_workbook(), tmp_path / "budget.xlsx")
    _store_result(path, "SUM(B2:B3)", "1250000")

    result = normalize(path, DIGEST)

    total = next(segment for segment in result.segments if segment.locator.get("cell") == "B4")
    assert total.locator["formula"] == "=SUM(B2:B3)"
    assert "computed_value" not in total.locator
    # Plain cells gain no formula key, and the table carries the result, not the formula.
    assert _by_locator(result)[_key(sheet="Budget", cell="B2")] == "750000"
    assert result.tables[0].rows[-1] == ["Total", 1250000]


def test_computed_value_is_found_by_search_with_its_formula(tmp_path: Path) -> None:
    path = _save_workbook(_budget_workbook(), tmp_path / "budget.xlsx")
    _store_result(path, "SUM(B2:B3)", "1250000")

    hits = _search(path, tmp_path / "home", "1250000")

    assert [hit["citation"]["locator"] for hit in hits] == [  # pyright: ignore[reportIndexIssue]
        {"sheet": "Budget", "cell": "B4", "formula": "=SUM(B2:B3)"}
    ]


def test_formula_without_a_stored_result_indexes_its_text_and_says_so(tmp_path: Path) -> None:
    path = _save_workbook(_budget_workbook(), tmp_path / "never-calculated.xlsx")

    located = _by_locator(normalize(path, DIGEST))

    key = _key(sheet="Budget", cell="B4", formula="=SUM(B2:B3)", computed_value="not saved in file")
    assert located[key] == "=SUM(B2:B3)"
    assert "1250000" not in located.values()


def test_hidden_sheet_content_is_retrievable_and_marked_hidden(tmp_path: Path) -> None:
    workbook = Workbook()
    visible = workbook.active
    assert visible is not None
    visible.title = "Summary"
    visible["A1"] = "Nothing to see"
    hidden = workbook.create_sheet("Workings")
    hidden.sheet_state = "hidden"
    hidden["C3"] = "ARCHIV-HIDDEN-SHEET-MARKER"
    path = _save_workbook(workbook, tmp_path / "hidden.xlsx")

    result = normalize(path, DIGEST)
    hits = _search(path, tmp_path / "home", "ARCHIV-HIDDEN-SHEET-MARKER")

    assert _by_locator(result)[_key(sheet="Summary", cell="A1")] == "Nothing to see"
    assert [hit["citation"]["locator"] for hit in hits] == [  # pyright: ignore[reportIndexIssue]
        {"sheet": "Workings", "cell": "C3", "hidden": "sheet"}
    ]
    assert result.metadata["hidden_sheets"] == ["Workings"]


def test_hidden_rows_and_columns_are_marked(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Grid"
    sheet["A1"] = "shown"
    sheet["A2"] = "row hidden"
    sheet["B1"] = "column hidden"
    sheet["B2"] = "both hidden"
    sheet.row_dimensions[2].hidden = True
    sheet.column_dimensions["B"].hidden = True
    path = _save_workbook(workbook, tmp_path / "grid.xlsx")

    located = _by_locator(normalize(path, DIGEST))

    assert located[_key(sheet="Grid", cell="A1")] == "shown"
    assert located[_key(sheet="Grid", cell="A2", hidden="row")] == "row hidden"
    assert located[_key(sheet="Grid", cell="B1", hidden="column")] == "column hidden"
    assert located[_key(sheet="Grid", cell="B2", hidden="row and column")] == "both hidden"


def test_merged_header_reaches_its_whole_span(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Quarters"
    sheet["A1"] = "Revenue"
    sheet.merge_cells("A1:C1")
    for column, amount in zip("ABC", (10, 20, 30), strict=True):
        sheet[f"{column}2"] = amount
    path = _save_workbook(workbook, tmp_path / "merged.xlsx")

    result = normalize(path, DIGEST)

    located = _by_locator(result)
    assert located[_key(sheet="Quarters", cell="A1", merged="A1:C1")] == "Revenue"
    # One segment for the header, not one per covered cell.
    assert list(located.values()).count("Revenue") == 1
    assert result.tables[0].rows == [["Revenue", "Revenue", "Revenue"], [10, 20, 30]]


def test_merge_across_a_whole_row_is_clipped_to_the_cells_in_use(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "Banner"
    sheet["A2"] = "x"
    sheet["B2"] = "y"
    sheet.merge_cells("A1:XFD1")
    path = _save_workbook(workbook, tmp_path / "banner.xlsx")

    result = normalize(path, DIGEST)

    assert result.tables[0].rows == [["Banner", "Banner"], ["x", "y"]]


def test_xlsx_hyperlink_targets_are_preserved(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Links"
    sheet["A1"] = "Supplier site"
    sheet["A1"].hyperlink = "https://example.org/supplier"
    sheet["A2"] = "Jump to totals"
    sheet["A2"].hyperlink = "#Links!B9"
    path = _save_workbook(workbook, tmp_path / "links.xlsx")

    located = _by_locator(normalize(path, DIGEST))

    assert located[_key(sheet="Links", cell="A1")] == "Supplier site"
    assert located[_key(sheet="Links", cell="A1", hyperlink=1)] == "https://example.org/supplier"
    assert located[_key(sheet="Links", cell="A2", hyperlink=1)] == "#Links!B9"


def test_xlsx_hyperlink_with_a_missing_relationship_is_skipped(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "Broken link"
    sheet["A1"].hyperlink = "https://example.org/gone"
    path = _save_workbook(workbook, tmp_path / "broken.xlsx")
    _rewrite(
        path,
        {},
        {
            "xl/worksheets/_rels/sheet1.xml.rels": lambda xml: re.sub(
                r"<Relationship [^>]*hyperlink[^>]*/>", "", xml
            )
        },
    )

    result = normalize(path, DIGEST)

    assert [segment.text for segment in result.segments] == ["Broken link"]


def test_xlsx_chart_titles_and_series_names_are_extracted(tmp_path: Path) -> None:
    workbook = _budget_workbook()
    sheet = workbook.active
    assert sheet is not None
    chart = BarChart()
    chart.title = "Spending by category"
    chart.y_axis.title = "Euros"
    chart.add_data(Reference(sheet, min_col=2, min_row=2, max_row=3))  # pyright: ignore[reportUnknownMemberType]
    chart.anchor = "D2"
    sheet.add_chart(chart)  # pyright: ignore[reportUnknownMemberType]
    path = _save_workbook(workbook, tmp_path / "chart.xlsx")
    _rewrite(
        path,
        {},
        {
            # openpyxl names a series only by reference; an application also saves
            # the name it last read from that reference.
            "xl/charts/chart1.xml": lambda xml: xml.replace(
                "<ser><idx",
                "<ser><tx><strRef><f>'Budget'!A1</f><strCache><ptCount val=\"1\"/>"
                '<pt idx="0"><v>Annual costs</v></pt></strCache></strRef></tx><idx',
                1,
            )
        },
    )

    result = normalize(path, DIGEST)

    located = _by_locator(result)
    assert located[_key(sheet="Budget", chart=1)] == "Spending by category\nEuros\nAnnual costs"
    assert result.metadata["charts"] == 1


def test_xlsx_entity_declaration_refuses_the_workbook(tmp_path: Path) -> None:
    path = _save_workbook(_budget_workbook(), tmp_path / "entity.xlsx")
    _rewrite(
        path,
        {},
        {
            XLSX_SHEET: lambda xml: (
                '<?xml version="1.0"?><!DOCTYPE worksheet [<!ENTITY boom "boom">]>'
                + xml.split("?>", 1)[-1]
            )
        },
    )

    with pytest.raises(MalformedInputError, match="declarations and entities"):
        normalize(path, DIGEST)


def test_xlsx_declared_dimension_cannot_force_unbounded_iteration(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(normalize_xlsx_module, "MAX_CELL_POSITIONS", 1_000)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "top"
    path = _save_workbook(workbook, tmp_path / "dimension.xlsx")
    _rewrite(
        path,
        {},
        {
            XLSX_SHEET: lambda xml: re.sub(
                r'<dimension ref="[^"]*"/>', '<dimension ref="A1:XFD1048576"/>', xml
            )
        },
    )

    with pytest.raises(MalformedInputError, match="cell position limit"):
        normalize(path, DIGEST)


def test_xlsx_chart_sheet_and_very_hidden_sheet_are_read(tmp_path: Path) -> None:
    workbook = Workbook()
    data = workbook.active
    assert data is not None
    data.title = "Data"
    data["A1"] = 1
    chart_sheet = workbook.create_chartsheet("Plot")
    chart = BarChart()
    chart.title = "Plotted on its own sheet"
    chart.add_data(Reference(data, min_col=1, min_row=1, max_row=1))  # pyright: ignore[reportUnknownMemberType]
    chart_sheet.add_chart(chart)  # pyright: ignore[reportUnknownMemberType]
    secret = workbook.create_sheet("Secret")
    secret.sheet_state = "veryHidden"
    secret["A1"] = "deep"
    path = _save_workbook(workbook, tmp_path / "chartsheet.xlsx")

    located = _by_locator(normalize(path, DIGEST))

    assert located[_key(sheet="Plot", chart=1)] == "Plotted on its own sheet"
    assert located[_key(sheet="Secret", cell="A1", hidden="sheet")] == "deep"
