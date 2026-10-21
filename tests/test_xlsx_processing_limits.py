from __future__ import annotations

import os
import re
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from zipfile import ZIP_STORED, ZipFile

import pytest
from openpyxl import Workbook

from archiv.hashing import sha256_file
from archiv.ingestion import ingest_file, xlsx_process
from archiv.ingestion.limits import NativeResourceLimitError
from archiv.ingestion.normalizers import MalformedInputError, normalize
from archiv.storage.database import ArchivDatabase
from archiv.storage.layout import ArchivLayout


def _workbook(path: Path, count: int = 1) -> Path:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    for row in range(1, count + 1):
        sheet.cell(row, 1, f"ARCHIV-WORKER-{row}")
    book.save(path)
    return path


def _instrument_library(monkeypatch: pytest.MonkeyPatch, marker: Path, *, stall: bool) -> None:
    # Inject into the disposable child, not the parent. A parent-only monkeypatch
    # would test nothing after isolation moved the actual library call elsewhere.
    hook = (
        "import openpyxl, os, time\n"
        "from pathlib import Path\n"
        "def instrumented_load(*args, **kwargs):\n"
        f"    Path({str(marker)!r}).write_text(str(os.getpid()))\n"
        + ("    time.sleep(60)\n" if stall else "    raise AssertionError('library was called')\n")
        + "openpyxl.load_workbook = instrumented_load\n"
    )
    bootstrap = xlsx_process._WORKER_BOOTSTRAP  # pyright: ignore[reportPrivateUsage]
    monkeypatch.setattr(
        xlsx_process,
        "_WORKER_BOOTSTRAP",
        bootstrap.replace(
            "from archiv.ingestion.xlsx_process import worker_main",
            hook + "from archiv.ingestion.xlsx_process import worker_main",
        ),
    )


def test_repeated_sheet_references_cannot_multiply_unbounded_work(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = _workbook(tmp_path / "repeated.xlsx")
    with ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    workbook = members["xl/workbook.xml"].decode()
    entry = re.search(r"<sheet [^>]*/>", workbook)
    assert entry is not None
    copies = "".join(
        entry[0]
        .replace('name="Sheet"', f'name="Copy {i}"')
        .replace('sheetId="1"', f'sheetId="{i + 1}"')
        for i in range(1, 200)
    )
    members["xl/workbook.xml"] = workbook.replace(entry[0], entry[0] + copies).encode()
    with ZipFile(path, "w", ZIP_STORED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    marker = tmp_path / "library-called"
    _instrument_library(monkeypatch, marker, stall=False)
    digest = sha256_file(path)

    with pytest.raises(MalformedInputError, match="repeated worksheet part"):
        normalize(path, digest)

    assert not marker.exists()
    assert sha256_file(path) == digest


def test_processing_budget_covers_library_calls(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = _workbook(tmp_path / "stall.xlsx")
    marker = tmp_path / "worker-pid"
    _instrument_library(monkeypatch, marker, stall=True)
    monkeypatch.setattr(xlsx_process, "MAX_TIMEOUT_SECONDS", 2)

    with pytest.raises(NativeResourceLimitError, match="processing time limit"):
        normalize(path, sha256_file(path))

    assert marker.exists(), "the test must reach the actual library call before timing out"
    with pytest.raises(ProcessLookupError):
        os.kill(int(marker.read_text()), 0)


def test_budget_expiry_preserves_source_and_records_failure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = _workbook(tmp_path / "refused.xlsx")
    digest = sha256_file(path)
    marker = tmp_path / "worker-pid"
    _instrument_library(monkeypatch, marker, stall=True)
    monkeypatch.setattr(xlsx_process, "MAX_TIMEOUT_SECONDS", 2)
    layout = ArchivLayout.resolve(tmp_path / "home")
    layout.ensure()
    database = ArchivDatabase(layout.database)
    database.initialize()

    with pytest.raises(NativeResourceLimitError, match="processing time limit"):
        ingest_file(path, home=layout.root)

    assert marker.exists()
    assert sha256_file(path) == digest
    assert not layout.original_path(digest).exists()
    assert not layout.derived_root(digest).exists()
    with database.connect() as connection:
        assert connection.execute("SELECT COUNT(*) FROM objects").fetchone()[0] == 0
        failure = connection.execute("SELECT error FROM ingestion_failures").fetchone()
    assert failure is not None and "processing time limit" in failure[0]


def test_ordinary_workbook_still_ingests_within_budget(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "ordinary.xlsx", count=10_000)
    digest = sha256_file(path)
    home = tmp_path / "home"

    result = ingest_file(path, home=home)

    assert result.source_hash_unchanged
    assert sha256_file(path) == digest
    assert sha256_file(ArchivLayout.resolve(home).original_path(digest)) == digest


def test_normalization_from_a_worker_thread_is_supported(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "threaded.xlsx")
    digest = sha256_file(path)
    with ThreadPoolExecutor(max_workers=1) as executor:
        result = executor.submit(normalize, path, digest).result(timeout=20)
    assert [segment.text for segment in result.segments] == ["ARCHIV-WORKER-1"]


@pytest.mark.parametrize(
    "bootstrap",
    (
        "import os,signal; os.kill(os.getpid(), signal.SIGKILL)",
        "raise MemoryError('generated worker failure')",
    ),
)
def test_resource_terminated_workers_never_return_a_document(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, bootstrap: str
) -> None:
    path = _workbook(tmp_path / "terminated.xlsx")
    monkeypatch.setattr(xlsx_process, "_WORKER_BOOTSTRAP", bootstrap)
    # A Python startup failure and a signal death both fail closed; no successful
    # document is returned even when the worker never produced a response.
    with pytest.raises((NativeResourceLimitError, MalformedInputError)):
        normalize(path, sha256_file(path))


@pytest.mark.parametrize("response", ("not-json", "[]", '{"status":"succeeded","document":{}}'))
def test_malformed_worker_response_fails_closed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, response: str
) -> None:
    path = _workbook(tmp_path / "bad-response.xlsx")
    monkeypatch.setattr(xlsx_process, "_WORKER_BOOTSTRAP", f"print({response!r})")
    with pytest.raises(MalformedInputError):
        normalize(path, sha256_file(path))


@pytest.mark.skipif(not Path("/proc/self/fd").is_dir(), reason="descriptor accounting needs procfs")
def test_timeout_does_not_leak_worker_descriptors(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = _workbook(tmp_path / "descriptors.xlsx")
    marker = tmp_path / "worker-pid"
    _instrument_library(monkeypatch, marker, stall=True)
    monkeypatch.setattr(xlsx_process, "MAX_TIMEOUT_SECONDS", 2)
    before = len(list(Path("/proc/self/fd").iterdir()))

    with pytest.raises(NativeResourceLimitError):
        normalize(path, sha256_file(path))

    assert marker.exists()
    assert len(list(Path("/proc/self/fd").iterdir())) <= before


@pytest.mark.parametrize("shape", ("decoy", "no-leading-slash", "duplicate-first-wins"))
def test_manifest_and_relationship_workbook_parts_must_agree(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, shape: str
) -> None:
    # openpyxl picks the workbook part from [Content_Types].xml, using the first
    # matching Override and PartName[1:] literally; this reader picks it from the
    # package relationships. A file where those disagree would let the budgets
    # describe one workbook while openpyxl parses another.
    path = _workbook(tmp_path / "decoy.xlsx")
    with ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members["xl/decoy.xml"] = members["xl/workbook.xml"]
    content_types = members["[Content_Types].xml"].decode("utf-8")
    found = re.search(r'<Override PartName="/xl/workbook\.xml"[^>]*/>', content_types)
    assert found is not None
    element = found.group(0)
    if shape == "decoy":
        content_types = content_types.replace(
            element, element.replace("/xl/workbook.xml", "/xl/decoy.xml")
        )
    elif shape == "no-leading-slash":
        content_types = content_types.replace(
            element, element.replace('PartName="/xl/workbook.xml"', 'PartName="xl/workbook.xml"')
        )
    else:
        # openpyxl takes the first matching Override. A reader that keys Overrides
        # by part name would let a later duplicate of the decoy part win, keep the
        # real workbook type, and agree while openpyxl still parses the decoy.
        workbook_type = re.search(r'ContentType="([^"]+)"', element)
        assert workbook_type is not None
        decoy = element.replace("/xl/workbook.xml", "/xl/decoy.xml")
        overwritten = decoy.replace(
            workbook_type.group(1),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml",
        )
        content_types = content_types.replace(element, decoy + overwritten + element)
    members["[Content_Types].xml"] = content_types.encode("utf-8")
    with ZipFile(path, "w", ZIP_STORED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    marker = tmp_path / "library-called"
    _instrument_library(monkeypatch, marker, stall=False)

    with pytest.raises(MalformedInputError, match="disagrees with the package manifest"):
        normalize(path, sha256_file(path))

    assert not marker.exists()


def test_non_utf8_stdout_locale_cannot_corrupt_the_response(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "unicode.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet["A1"] = "日本語 😀 ARCHIV"
    book.save(path)
    bootstrap = xlsx_process._WORKER_BOOTSTRAP  # pyright: ignore[reportPrivateUsage]
    monkeypatch.setattr(
        xlsx_process,
        "_WORKER_BOOTSTRAP",
        'import sys\nsys.stdout.reconfigure(encoding="latin-1")\n' + bootstrap,
    )

    result = normalize(path, sha256_file(path))

    assert [segment.text for segment in result.segments] == ["日本語 😀 ARCHIV"]


def test_worker_refuses_a_file_that_no_longer_matches_its_hash(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "changed.xlsx")

    with pytest.raises(MalformedInputError, match="changed before the worker read it"):
        normalize(path, "0" * 64)


def test_oversized_response_is_reported_as_a_resource_limit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = _workbook(tmp_path / "oversized.xlsx")
    monkeypatch.setattr(
        xlsx_process,
        "_WORKER_BOOTSTRAP",
        "import os, resource\n"
        "resource.setrlimit(resource.RLIMIT_FSIZE, (16, 16))\n"
        "while True:\n"
        '    os.write(1, b"x" * 64)\n',
    )

    with pytest.raises(NativeResourceLimitError, match="could not write its complete result"):
        normalize(path, sha256_file(path))
