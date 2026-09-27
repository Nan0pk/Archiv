"""DOCX normalization: body, tables, headers, footers, notes, comments, and links.

``python-docx``'s ``document.paragraphs`` only lists paragraphs that sit directly in
the document body. Table cells, headers, footers, footnotes, endnotes and comments
live elsewhere in the package, so reading paragraphs alone silently drops them. This
module walks each of those parts explicitly and gives every recovered piece of text a
locator that says where in the document it came from.

``python-docx`` still opens and validates the package and resolves sections, styles
and relationships. The XML of each part is then read with the standard library, and
paragraph text follows exactly the rules ``python-docx`` uses for ``Paragraph.text``,
so body text is unchanged from the paragraphs-only reader.

Body paragraph numbering is also unchanged: ``paragraph`` is the 1-based position
among paragraphs directly in the body, counting empty ones, so "paragraph N" still
names the same text it did under the paragraphs-only reader. A citation saved before
this reader existed also records the whole derived document's hash and the segment's
position, so once a document is rebuilt with this reader those saved citations stop
validating; they fail closed rather than point at the wrong text.
"""

from __future__ import annotations

import re
from pathlib import Path
from xml.etree import ElementTree
from xml.parsers import expat

from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.part import Part

from archiv.contracts import NormalizedDocument, NormalizedSegment, NormalizedTable

__all__ = ["MAX_TABLE_CELLS", "MAX_TABLE_COLUMNS", "MAX_XML_NODES", "normalize_docx"]

# Word documents are bounded by the ingestion size limit, but a small package can
# still declare a very large table: one cell may claim to span any number of grid
# columns. The budget therefore counts grid positions, including spans, leading gaps
# and the padding that makes every row as wide as the widest, not just cells. Word
# itself allows 63 columns; the column ceiling leaves room for other producers.
MAX_TABLE_CELLS = 200_000
MAX_TABLE_COLUMNS = 1_000
# Real Word parts never declare a DTD. Refusing declarations outright rules out entity
# expansion, and the node ceiling bounds the parsed tree, as parse_odf_xml does for
# OpenDocument. The ceiling is higher than OpenDocument's because every run, text
# node and property of a long Word document is its own element.
MAX_XML_NODES = 5_000_000

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

_HEADING_STYLE = re.compile(r"^Heading ([1-9])$")
_NOTE_ID = re.compile(r"-?[0-9]{1,9}")
_NOTE_SKIP_TYPES = {"separator", "continuationSeparator", "continuationNotice"}
_HEADER_FOOTER_KINDS = (
    ("default", "header", "footer"),
    ("first", "first_page_header", "first_page_footer"),
    ("even", "even_page_header", "even_page_footer"),
)

type Locator = dict[str, object]


def _w(name: str) -> str:
    return f"{{{W}}}{name}"


def _attr(element: ElementTree.Element, name: str, default: str = "") -> str:
    return element.attrib.get(_w(name), default)


class _CellBudget:
    def __init__(self) -> None:
        self.used = 0

    def spend(self, positions: int = 1) -> None:
        self.used += positions
        if self.used > MAX_TABLE_CELLS:
            raise ValueError("DOCX table cell limit exceeded")


def _run_text(run: ElementTree.Element) -> str:
    """The same text ``python-docx`` gives for a ``w:r`` element."""

    parts: list[str] = []
    for child in run:
        if child.tag == _w("t"):
            parts.append(child.text or "")
        elif child.tag in {_w("tab"), _w("ptab")}:
            parts.append("\t")
        elif child.tag == _w("cr"):
            parts.append("\n")
        elif child.tag == _w("br"):
            # Only a line break is text; page and column breaks are layout.
            if _attr(child, "type", "textWrapping") == "textWrapping":
                parts.append("\n")
        elif child.tag == _w("noBreakHyphen"):
            parts.append("-")
    return "".join(parts)


def _paragraph_text(paragraph: ElementTree.Element) -> str:
    """The same text ``python-docx`` gives for ``Paragraph.text``.

    That is the direct runs of the paragraph plus the runs directly inside its
    hyperlinks, in order. Text inside tracked insertions, content controls and fields
    that are not plain runs is not included, exactly as before this reader existed.
    """

    parts: list[str] = []
    for child in paragraph:
        if child.tag == _w("r"):
            parts.append(_run_text(child))
        elif child.tag == _w("hyperlink"):
            parts.extend(_run_text(run) for run in child if run.tag == _w("r"))
    return "".join(parts)


def _hyperlink_targets(element: ElementTree.Element, part: Part) -> list[str]:
    """Link targets inside ``element``, in document order.

    An external link's address lives in the part's relationships, keyed by ``r:id``;
    an internal link names a bookmark in ``w:anchor``. A link whose relationship is
    missing from the package has no target to report, so it is skipped rather than
    guessed.
    """

    targets: list[str] = []
    for link in element.iter(_w("hyperlink")):
        target = ""
        relationship_id = link.attrib.get(f"{{{R}}}id")
        if relationship_id:
            relationship = part.rels.get(relationship_id)
            if relationship is not None:
                target = relationship.target_ref
        anchor = _attr(link, "anchor")
        if anchor:
            target = f"{target}#{anchor}"
        if target:
            targets.append(target)
    return targets


def _link_segments(
    element: ElementTree.Element, part: Part, locator: Locator
) -> list[NormalizedSegment]:
    return [
        NormalizedSegment(locator={**locator, "hyperlink": number}, text=target)
        for number, target in enumerate(_hyperlink_targets(element, part), 1)
    ]


def _is_vertical_continuation(cell: ElementTree.Element) -> bool:
    properties = cell.find(_w("tcPr"))
    if properties is None:
        return False
    merge = properties.find(_w("vMerge"))
    if merge is None:
        return False
    return _attr(merge, "val", "continue") == "continue"


def _grid_value(properties: ElementTree.Element | None, name: str) -> int:
    if properties is None:
        return 0
    child = properties.find(_w(name))
    if child is None:
        return 0
    raw = _attr(child, "val", "0")
    try:
        value = int(raw)
    except ValueError as error:
        raise ValueError(f"DOCX table has an invalid {name} value") from error
    if value < 0 or value > MAX_TABLE_COLUMNS:
        raise ValueError(f"DOCX table has an out-of-range {name} value")
    return value


def _table(
    table: ElementTree.Element,
    part: Part,
    prefix: Locator,
    budget: _CellBudget,
) -> tuple[list[NormalizedSegment], NormalizedTable | None]:
    """One segment per non-empty cell, located by table, row and grid column.

    Cell text joins the cell's non-empty paragraphs with newlines, including any table
    nested inside the cell, so nothing in a nested table is dropped. A horizontally
    merged cell is located at its first grid column; the cells a vertical merge covers
    carry no text of their own and yield nothing.
    """

    segments: list[NormalizedSegment] = []
    rows: list[list[object | None]] = []
    for row_number, row in enumerate(table.iterfind(_w("tr")), 1):
        column = _grid_value(row.find(_w("trPr")), "gridBefore") + 1
        budget.spend(column - 1)
        values: list[object | None] = [None] * (column - 1)
        for cell in row.iterfind(_w("tc")):
            span = max(_grid_value(cell.find(_w("tcPr")), "gridSpan"), 1)
            if column - 1 + span > MAX_TABLE_COLUMNS:
                raise ValueError("DOCX table column limit exceeded")
            budget.spend(span)
            text: str | None = None
            if not _is_vertical_continuation(cell):
                paragraphs = [_paragraph_text(p) for p in cell.iter(_w("p"))]
                joined = "\n".join(value for value in paragraphs if value)
                if joined:
                    text = joined
                    locator: Locator = {**prefix, "row": row_number, "column": column}
                    segments.append(NormalizedSegment(locator=locator, text=joined))
                    segments.extend(_link_segments(cell, part, locator))
            values.append(text)
            values.extend([None] * (span - 1))
            column += span
        rows.append(values)
    if not any(value is not None for row in rows for value in row):
        return segments, None
    width = max(len(row) for row in rows)
    budget.spend(sum(width - len(row) for row in rows))
    padded = [row + [None] * (width - len(row)) for row in rows]
    return segments, NormalizedTable(locator=dict(prefix), rows=padded)


def _heading_level(paragraph: ElementTree.Element, style_levels: dict[str, int]) -> int | None:
    """A heading level from the paragraph's own outline level, else from its style."""

    properties = paragraph.find(_w("pPr"))
    if properties is None:
        return None
    outline = properties.find(_w("outlineLvl"))
    if outline is not None:
        raw = _attr(outline, "val")
        # Outline levels 0-8 are headings 1-9; 9 means body text, even under a
        # heading style.
        if raw.isdigit() and int(raw) < 9:
            return int(raw) + 1
        if raw == "9":
            return None
    style = properties.find(_w("pStyle"))
    if style is not None:
        return style_levels.get(_attr(style, "val"))
    return None


def _container(
    root: ElementTree.Element,
    part: Part,
    prefix: Locator,
    budget: _CellBudget,
    *,
    style_levels: dict[str, int] | None = None,
) -> tuple[list[NormalizedSegment], list[NormalizedTable], int, int]:
    """Walk the paragraphs and tables directly inside one story, in document order.

    Returns the segments, the tables, and how many paragraphs and tables were seen,
    counting empty ones.
    """

    segments: list[NormalizedSegment] = []
    tables: list[NormalizedTable] = []
    paragraph_number = 0
    table_number = 0
    for child in root:
        if child.tag == _w("p"):
            paragraph_number += 1
            text = _paragraph_text(child)
            if not text:
                continue
            locator: Locator = {**prefix, "paragraph": paragraph_number}
            if style_levels is not None:
                level = _heading_level(child, style_levels)
                if level is not None:
                    locator["heading_level"] = level
            segments.append(NormalizedSegment(locator=locator, text=text))
            segments.extend(_link_segments(child, part, locator))
        elif child.tag == _w("tbl"):
            table_number += 1
            table_segments, table = _table(child, part, {**prefix, "table": table_number}, budget)
            segments.extend(table_segments)
            if table is not None:
                tables.append(table)
    return segments, tables, paragraph_number, table_number


class _DeclarationRefused(Exception):
    pass


def _refuse_declaration(*_: object) -> None:
    raise _DeclarationRefused


def _refuse_declarations(data: bytes) -> None:
    """Refuse any DTD or entity declaration, whatever the part's text encoding.

    A byte search for ``<!DOCTYPE`` misses a part encoded as UTF-16, which the parser
    still decodes and expands. So the check is made by the XML parser itself, which
    stops the moment a declaration starts, before any entity is expanded.
    """

    scanner = expat.ParserCreate()
    scanner.StartDoctypeDeclHandler = _refuse_declaration
    scanner.EntityDeclHandler = _refuse_declaration
    try:
        scanner.Parse(data, True)
    except _DeclarationRefused:
        raise ValueError("DOCX XML declarations and entities are not allowed") from None
    except expat.ExpatError as error:
        raise ValueError("DOCX package part is not well-formed XML") from error


def _part_root(part: Part) -> ElementTree.Element:
    data = part.blob
    _refuse_declarations(data)
    try:
        root = ElementTree.fromstring(data)
    except ElementTree.ParseError as error:
        raise ValueError("DOCX package part is not well-formed XML") from error
    if sum(1 for _ in root.iter()) > MAX_XML_NODES:
        raise ValueError("DOCX XML node limit exceeded")
    return root


def _style_heading_levels(document: DocxDocument) -> dict[str, int]:
    """Map each built-in heading style's identifier to its level (``Heading 2`` -> 2)."""

    levels: dict[str, int] = {}
    for style in document.styles:
        match = _HEADING_STYLE.match(style.name or "")
        if match and style.style_id:
            levels[style.style_id] = int(match.group(1))
    return levels


def _headers_and_footers(
    document: DocxDocument, budget: _CellBudget
) -> tuple[list[NormalizedSegment], list[NormalizedTable]]:
    """Each section's own header and footer stories.

    A header or footer that is linked to the previous section has no content of its
    own, so it is skipped rather than repeated under every section that shows it.
    """

    segments: list[NormalizedSegment] = []
    tables: list[NormalizedTable] = []
    for section_number, section in enumerate(document.sections, 1):
        for kind, header_attribute, footer_attribute in _HEADER_FOOTER_KINDS:
            for role, attribute in (("header", header_attribute), ("footer", footer_attribute)):
                story = getattr(section, attribute)
                if story.is_linked_to_previous:
                    continue
                part: Part = story.part
                story_segments, story_tables, _, _ = _container(
                    _part_root(part),
                    part,
                    {"section": section_number, role: kind},
                    budget,
                )
                segments.extend(story_segments)
                tables.extend(story_tables)
    return segments, tables


def _notes(
    document_part: Part,
    relationship_type: str,
    note_tag: str,
    key: str,
    budget: _CellBudget,
) -> tuple[list[NormalizedSegment], list[NormalizedTable]]:
    """Footnotes, endnotes or comments, each located by the note's own ``w:id``.

    The separator entries Word stores alongside real footnotes and endnotes are not
    content and are skipped.
    """

    try:
        part = document_part.part_related_by(relationship_type)
    except KeyError:
        return [], []
    segments: list[NormalizedSegment] = []
    tables: list[NormalizedTable] = []
    for note in _part_root(part).iterfind(_w(note_tag)):
        if _attr(note, "type") in _NOTE_SKIP_TYPES:
            continue
        raw_id = _attr(note, "id")
        # A malformed identifier is kept as written rather than refusing the document.
        note_id: object = int(raw_id) if _NOTE_ID.fullmatch(raw_id) else raw_id
        note_segments, note_tables, _, _ = _container(note, part, {key: note_id}, budget)
        segments.extend(note_segments)
        tables.extend(note_tables)
    return segments, tables


def normalize_docx(
    path: Path,
    digest: str,
    *,
    source_name: str,
    media_type: str,
) -> NormalizedDocument:
    document = Document(str(path))
    budget = _CellBudget()
    document_part = document.part

    body = _part_root(document_part).find(_w("body"))
    if body is None:
        raise ValueError("DOCX package has no document body")
    segments, tables, body_paragraphs, body_tables = _container(
        body,
        document_part,
        {},
        budget,
        style_levels=_style_heading_levels(document),
    )

    story_segments, story_tables = _headers_and_footers(document, budget)
    segments.extend(story_segments)
    tables.extend(story_tables)

    for relationship_type, note_tag, key in (
        (RT.FOOTNOTES, "footnote", "footnote"),
        (RT.ENDNOTES, "endnote", "endnote"),
        (RT.COMMENTS, "comment", "comment"),
    ):
        note_segments, note_tables = _notes(document_part, relationship_type, note_tag, key, budget)
        segments.extend(note_segments)
        tables.extend(note_tables)

    return NormalizedDocument(
        object_sha256=digest,
        media_type=media_type,
        kind="docx",
        source_name=source_name,
        segments=segments,
        tables=tables,
        metadata={
            "paragraphs": body_paragraphs,
            "tables": body_tables,
            "table_cells": budget.used,
        },
    )
